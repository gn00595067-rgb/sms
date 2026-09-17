"""
builder.py — 自訂條件報表產生器的引擎（CLAUDE_CODE_TASK_4 §5.3）

安全模型（同 reports/query.py）：
  - 維度 / 指標 / 篩選欄 全部白名單；使用者輸入只出現在 params（參數化綁定）。
  - 比率永遠事後算（SQL 只抓分子分母）；任何地方不 avg(margin)。
  - SALES 角色自動限縮成 salesperson_merged = 自己。
  - definition JSON 從 DB 載回後仍要經 validate_definition 白名單驗證再組 SQL。

來源：v_report_line（線層；集團毛利用 group_profit_alloc 分攤欄）。
"""
from __future__ import annotations

from datetime import date

import pandas as pd

from core import analysis as A

GRAIN = "v_report_line"

# key: (中文, SQL 運算式, 排序)  —— fixed:a,b,c 表示固定欄序
DIMENSIONS: dict[str, tuple] = {
    "year":              ("年", "perf_year::text", "asc"),
    "quarter":           ("季", "(perf_year::text || 'Q' || perf_quarter::text)", "asc"),
    "month":             ("月", "ym", "asc"),
    "company":           ("公司", "company", "fixed:聲活,東吳,鉑霖,瑞迪"),
    "report_platform":   ("報表平台", "report_platform", "fixed:全家企頻,萬家福,新鮮視,廣播,健康視,營運,其它"),
    "platform":          ("平台", "platform", "asc"),
    "media_channel":     ("電台/頻道", "media_channel", "asc"),
    "region":            ("區域", "region_text", "asc"),
    "salesperson":       ("業務", "salesperson_merged", "measure"),
    "business_group":    ("組別", "business_group", "asc"),
    "customer":          ("客戶", "customer", "measure"),
    "industry":          ("產業別", "coalesce(industry,'未分類')", "measure"),
    "customer_category": ("客戶類別", "customer_category", "measure"),
    "sales_category":    ("業績類別", "sales_category", "asc"),
    "line_type":         ("線類型", "line_type", "asc"),
    "is_barter":         ("交換", "case when is_barter then '交換' else '一般' end", "asc"),
    "contract_no":       ("合約編號", "contract_no", "asc"),
}

# key: (中文, SQL 彙總 or None, 格式, 比率分子分母)
MEASURES: dict[str, tuple] = {
    "line_count":     ("筆數", "count(*)", "int", None),
    "deal_count":     ("合約數", "count(distinct deal_id)", "int", None),
    "customers":      ("客戶數(不重複)", "count(distinct customer_id)", "int", None),
    "gross_amount":   ("實收", "sum(gross_amount)", "money", None),
    "net_amount":     ("除佣實收", "sum(net_amount)", "money", None),
    "cost_amount":    ("成本(實付)", "sum(cost_amount)", "money", None),
    "booked_profit":  ("帳上毛利", "sum(net_amount - cost_amount)", "money", None),
    "booked_margin":  ("帳上利率", None, "pct", ("booked_profit", "net_amount")),
    "group_profit":   ("集團毛利", "sum(group_profit_alloc)", "money", None),
    "group_margin":   ("集團利率", None, "pct", ("group_profit", "net_amount")),
    "recognized":     ("認定業績", "sum(recognized_amount)", "money", None),
    "avg_net":        ("平均單筆", None, "money", ("net_amount", "deal_count")),
    "purchased_slots": ("購買檔次", "sum(purchased_slots)", "int", None),
    "total_seconds":  ("總秒數", "sum(total_seconds)", "int", None),
}

# 篩選欄白名單：filter key -> SQL 欄運算式（值一律參數化）
FILTER_COLS: dict[str, str] = {
    "company": "company", "report_platform": "report_platform", "platform": "platform",
    "media_channel": "media_channel", "region": "region_text", "salesperson": "salesperson_merged",
    "business_group": "business_group", "industry": "coalesce(industry,'未分類')",
    "customer_category": "customer_category", "sales_category": "sales_category", "line_type": "line_type",
}

# 明細模式可輸出欄（白名單）
DETAIL_COLUMNS = ["ym", "contract_no", "company", "customer", "salesperson_merged",
                  "report_platform", "platform", "media_channel", "line_type", "air_period_text",
                  "gross_amount", "net_amount", "cost_amount", "booked_profit"]

MONEY_FMT = {"money", "int"}


def default_definition() -> dict:
    return {"v": 1, "name": "", "period": {"from": None, "to": None, "yoy": False},
            "scope": {"preset": "analysis"},
            "filters": {}, "rows": ["company"], "col": None,
            "measures": ["net_amount", "cost_amount", "booked_profit", "booked_margin"],
            "show": {"subtotal": False, "total": True, "share": True, "rank": False,
                     "top_n": None, "sort": ["net_amount", "desc"]},
            "detail_columns": None}


def validate_definition(defn: dict) -> dict:
    """白名單清洗：未知維度/指標/篩選欄略過，缺欄補預設。回傳乾淨 defn。"""
    d = default_definition()
    if not isinstance(defn, dict):
        return d
    d.update({k: v for k, v in defn.items() if k in d})
    d["rows"] = [k for k in (defn.get("rows") or []) if k in DIMENSIONS][:3]
    col = defn.get("col")
    d["col"] = col if col in DIMENSIONS else None
    d["measures"] = [k for k in (defn.get("measures") or []) if k in MEASURES] or ["net_amount"]
    d["filters"] = {k: v for k, v in (defn.get("filters") or {}).items()
                    if (k in FILTER_COLS or k in ("customer_like", "net_min", "margin_below")) and v not in (None, "", [])}
    d["scope"] = A.resolve_scope(defn.get("scope") or {"preset": "analysis"})
    show = dict(d["show"]); show.update(defn.get("show") or {})
    if show.get("sort") and (not isinstance(show["sort"], (list, tuple)) or len(show["sort"]) != 2
                             or show["sort"][0] not in (set(MEASURES) | set(DIMENSIONS))):
        show["sort"] = ["net_amount", "desc"]
    d["show"] = show
    dc = defn.get("detail_columns")
    d["detail_columns"] = [c for c in dc if c in DETAIL_COLUMNS] if dc else None
    return d


def _base_measure_keys(measures: list[str]) -> list[str]:
    """展開比率指標成需要的分子分母基礎指標（有 agg 的）。"""
    need: list[str] = []
    for m in measures:
        pair = MEASURES[m][3]
        if pair is None:
            need.append(m)
        else:
            need.extend(pair)
    # 去重、保序，且只留有 agg 的
    seen, out = set(), []
    for m in need + ["net_amount"]:      # 一律含 net_amount（給佔比）
        if m not in seen and MEASURES[m][1] is not None:
            seen.add(m); out.append(m)
    return out


def _where(defn: dict, user: dict | None) -> tuple[list[str], list]:
    where = ["perf_ym between %s and %s"] + A._scope_where(defn["scope"])
    p = defn["period"]
    params: list = [p.get("from") or date(2000, 1, 1), p.get("to") or date(2100, 1, 1)]
    for k, v in defn["filters"].items():
        if k in FILTER_COLS and isinstance(v, list) and v:
            where.append(f"{FILTER_COLS[k]} = any(%s)"); params.append(list(v))
        elif k == "customer_like" and v:
            pats = [f"%{x}%" for x in (v if isinstance(v, list) else [v])]
            where.append("customer ilike any(%s)"); params.append(pats)
    # SALES 限縮
    sc = A.sales_scope_name(user)
    if sc is not None:
        where.append("salesperson_merged = %s" if sc else "false")
        if sc:
            params.append(sc)
    return where, params


def _having(defn: dict) -> tuple[list[str], list]:
    having, params = [], []
    f = defn["filters"]
    if f.get("net_min") not in (None, ""):
        having.append("sum(net_amount) >= %s"); params.append(float(f["net_min"]))
    if f.get("margin_below") not in (None, ""):
        having.append("sum(net_amount) > 0 and (sum(net_amount)-sum(cost_amount)) < %s * sum(net_amount)")
        params.append(float(f["margin_below"]))
    return having, params


def build_sql(defn: dict, user: dict | None = None) -> tuple[str, list]:
    """回傳 (sql, params)。維度為空 → 明細模式（逐筆，白名單欄）。"""
    d = validate_definition(defn)
    dims = d["rows"] + ([d["col"]] if d["col"] else [])
    where, params = _where(d, user)

    if not d["rows"]:   # 明細模式
        cols = d["detail_columns"] or DETAIL_COLUMNS
        sql = f"select {', '.join(cols)} from {GRAIN}"
        if where:
            sql += " where " + " and ".join(where)
        sql += " order by perf_ym, contract_no limit 50000"
        return sql, params

    base = _base_measure_keys(d["measures"])
    sel = [f"{DIMENSIONS[k][1]} as {k}" for k in dims] + [f"{MEASURES[k][1]} as {k}" for k in base]
    sql = f"select {', '.join(sel)} from {GRAIN}"
    if where:
        sql += " where " + " and ".join(where)
    sql += " group by " + ", ".join(DIMENSIONS[k][1] for k in dims)
    having, hp = _having(d)
    if having:
        sql += " having " + " and ".join(having)
    params = params + hp
    return sql, params


def run(defn: dict, user: dict | None = None) -> pd.DataFrame:
    d = validate_definition(defn)
    sql, params = build_sql(d, user)
    df = A._df(sql, params)
    if df.empty:
        return df
    numcols = [c for c in df.columns if c in MEASURES and MEASURES[c][2] in MONEY_FMT or c in
               ("net_amount", "cost_amount", "booked_profit", "group_profit", "gross_amount",
                "recognized", "line_count", "deal_count", "customers", "purchased_slots", "total_seconds")]
    for c in numcols:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0)
    return df


def _fixed_order(dim: str) -> list[str] | None:
    s = DIMENSIONS[dim][2]
    return s.split("fixed:")[1].split(",") if s.startswith("fixed:") else None


def _measure_value(g: pd.DataFrame, m: str) -> float:
    """從已 group 的一段（含基礎欄）算出指標值（比率事後算）。"""
    agg, pair = MEASURES[m][1], MEASURES[m][3]
    if pair is None:
        return float(g[m].sum()) if m in g.columns else 0.0
    num, den = pair
    dn = float(g[den].sum()) if den in g.columns else 0.0
    nu = float(g[num].sum()) if num in g.columns else 0.0
    return (nu / dn) if dn else None


def shape(df: pd.DataFrame, defn: dict) -> tuple[pd.DataFrame, list]:
    """
    長表 → 顯示表 + 欄位規格 [(col, header, kind)]。處理樞紐（col 維度）、比率事後算、
    合計列、佔比、前 N、排序。明細模式（rows 空）直接回傳。
    """
    d = validate_definition(defn)
    if df is None or df.empty:
        return pd.DataFrame(), []
    if not d["rows"]:   # 明細模式
        cols = d["detail_columns"] or DETAIL_COLUMNS
        spec = [(c, MEASURES[c][0] if c in MEASURES else c,
                 MEASURES[c][2] if c in MEASURES else "text") for c in cols]
        return df, spec

    rows, col, measures, show = d["rows"], d["col"], d["measures"], d["show"]
    grand_net = float(df["net_amount"].sum()) if "net_amount" in df.columns else 0.0

    # 欄維度值（固定欄序優先，否則依金額）
    col_vals: list = []
    if col:
        present = list(dict.fromkeys(df[col].dropna().tolist()))
        fixed = _fixed_order(col)
        col_vals = ([c for c in fixed if c in present] if fixed
                    else sorted(present, key=lambda v: -float(df[df[col] == v]["net_amount"].sum())))

    spec: list = [(k, DIMENSIONS[k][0], "text") for k in rows]
    if col:
        for cv in col_vals:
            for m in measures:
                spec.append((f"{cv}｜{m}", f"{cv}｜{MEASURES[m][0]}", MEASURES[m][2]))
        for m in measures:   # 每列 合計欄群（跨欄維度）
            spec.append((f"合計｜{m}", f"合計｜{MEASURES[m][0]}", MEASURES[m][2]))
    else:
        for m in measures:
            spec.append((m, MEASURES[m][0], MEASURES[m][2]))
    if show.get("share"):
        spec.append(("__share", "佔比", "progress"))

    out_rows = []
    for keys, g in df.groupby(rows, dropna=False):
        keys = keys if isinstance(keys, tuple) else (keys,)
        row = {rows[i]: keys[i] for i in range(len(rows))}
        row["__sort"] = _measure_value(g, show.get("sort", ["net_amount"])[0]) or float(g["net_amount"].sum())
        if col:
            for cv in col_vals:
                sub = g[g[col] == cv]
                for m in measures:
                    row[f"{cv}｜{m}"] = _measure_value(sub, m)
            for m in measures:
                row[f"合計｜{m}"] = _measure_value(g, m)
        else:
            for m in measures:
                row[m] = _measure_value(g, m)
        row["__share"] = (float(g["net_amount"].sum()) / grand_net) if grand_net else 0.0
        out_rows.append(row)

    res = pd.DataFrame(out_rows)
    skey, sdir = show.get("sort", ["net_amount", "desc"])
    if "__sort" in res.columns:
        res = res.sort_values("__sort", ascending=(sdir != "desc"))
    if show.get("top_n"):
        res = res.groupby(rows[0], dropna=False, sort=False).head(int(show["top_n"]))
    res = res.reset_index(drop=True)

    if show.get("total") and len(res):
        tot = {rows[0]: "合計"}
        for c, _h, kind in spec:
            if c in (rows[0],) or c.startswith("__"):
                continue
            base_m = c.split("｜")[-1] if "｜" in c else c
            if base_m in MEASURES and MEASURES[base_m][3] is not None:
                tot[c] = _ratio_total(df, res, c, base_m, col)
            elif kind in MONEY_FMT:
                tot[c] = float(pd.to_numeric(res[c], errors="coerce").fillna(0).sum())
        tot["__share"] = 1.0
        res = pd.concat([res, pd.DataFrame([tot])], ignore_index=True)

    if "__sort" in res.columns:
        res = res.drop(columns="__sort")
    return res, spec


# ------------------------------------------------------------------ 儲存 / 載入（saved_report）
def list_saved(user: dict | None) -> list[dict]:
    """回傳可見的已存報表：自己的 + 共用的。SALES 也只看自己 + 共用。"""
    uid = (user or {}).get("id")
    rows = A._df(
        "select id, name, owner_id, is_shared, definition, updated_at, updated_by "
        "from saved_report where owner_id = %s or is_shared order by is_shared, name",
        (uid,))
    return rows.to_dict("records") if not rows.empty else []


def save_saved(name: str, defn: dict, user: dict, is_shared: bool = False,
               report_id: int | None = None) -> int:
    """新增或更新一個定義（白名單清洗後存）。SALES 不能分享。回傳 id。"""
    import json
    from db import transaction
    d = validate_definition(defn)
    d["name"] = name
    # date 不能直接進 jsonb → 轉字串
    if d["period"].get("from"):
        d["period"]["from"] = str(d["period"]["from"])
    if d["period"].get("to"):
        d["period"]["to"] = str(d["period"]["to"])
    if user.get("role") == "SALES":
        is_shared = False
    payload = json.dumps(d, ensure_ascii=False, default=str)
    uname = user.get("username")
    with transaction(uname) as cur:
        if report_id:
            cur.execute(
                "update saved_report set name=%s, definition=%s::jsonb, is_shared=%s, updated_by=%s "
                "where id=%s and owner_id=%s returning id",
                (name, payload, is_shared, uname, report_id, user["id"]))
        else:
            cur.execute(
                "insert into saved_report(name, owner_id, is_shared, definition, updated_by) "
                "values (%s,%s,%s,%s::jsonb,%s) "
                "on conflict (owner_id, name) do update set definition=excluded.definition, "
                "is_shared=excluded.is_shared, updated_by=excluded.updated_by returning id",
                (name, user["id"], is_shared, payload, uname))
        r = cur.fetchone()
    return int(r["id"]) if r else 0


def delete_saved(report_id: int, user: dict) -> None:
    from db import transaction
    with transaction(user.get("username")) as cur:
        cur.execute("delete from saved_report where id=%s and owner_id=%s", (report_id, user["id"]))


def load_definition(raw: dict) -> dict:
    """把存下來的定義（date 是字串）還原成可執行 defn（period 轉回 date）。"""
    from datetime import datetime
    d = validate_definition(raw)
    for k in ("from", "to"):
        v = (raw.get("period") or {}).get(k)
        if isinstance(v, str) and v:
            try:
                d["period"][k] = datetime.strptime(v[:10], "%Y-%m-%d").date()
            except ValueError:
                pass
    return d


# ------------------------------------------------------------------ Excel（多層表頭）
def build_excel(res: pd.DataFrame, spec: list, title: str, sub_lines: list[str]) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    FONT = "微軟正黑體"
    hfill = PatternFill("solid", fgColor="1F5F8B")
    hfont = Font(name=FONT, bold=True, color="FFFFFF")
    tfont = Font(name=FONT, bold=True, size=13)
    sfont = Font(name=FONT, size=9, color="666666")
    base = Font(name=FONT); bold = Font(name=FONT, bold=True)
    topb = Border(top=Side(style="thin"))
    wb = Workbook(); ws = wb.active; ws.title = "自訂報表"

    cols = [s[0] for s in spec]
    headers = [s[1] for s in spec]
    kinds = {s[0]: s[2] for s in spec}
    ncol = max(len(cols), 1)
    ws.cell(1, 1, title).font = tfont
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    for i, line in enumerate(sub_lines):
        ws.cell(2 + i, 1, line).font = sfont
        ws.merge_cells(start_row=2 + i, start_column=1, end_row=2 + i, end_column=ncol)

    # 兩層表頭：上層 = ｜前段（欄維度值），下層 = ｜後段（指標）
    hr_top = 2 + len(sub_lines) + 1
    hr_bot = hr_top + 1
    for j, (c, h) in enumerate(zip(cols, headers), 1):
        if "｜" in h:
            top, bot = h.split("｜", 1)
        else:
            top, bot = "", h
        tc = ws.cell(hr_top, j, top); tc.fill = hfill; tc.font = hfont; tc.alignment = Alignment(horizontal="center")
        bc = ws.cell(hr_bot, j, bot); bc.fill = hfill; bc.font = hfont; bc.alignment = Alignment(horizontal="center")

    r = hr_bot + 1
    n = len(res)
    for i in range(n):
        row = res.iloc[i]
        is_tot = str(row.get(cols[0])) == "合計"
        for j, c in enumerate(cols, 1):
            v = row.get(c)
            cell = ws.cell(r, j)
            if isinstance(v, float) and pd.isna(v):
                v = None
            cell.value = v
            cell.font = bold if is_tot else base
            if is_tot:
                cell.border = topb
            k = kinds.get(c)
            if v is not None and k in ("money", "int"):
                try:
                    cell.value = float(v); cell.number_format = "#,##0"; cell.alignment = Alignment(horizontal="right")
                except (ValueError, TypeError):
                    pass
            elif v is not None and k in ("pct", "progress"):
                try:
                    cell.value = float(v); cell.number_format = "0.0%"; cell.alignment = Alignment(horizontal="right")
                except (ValueError, TypeError):
                    pass
        r += 1

    ws.freeze_panes = ws.cell(hr_bot + 1, 2)
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(20, len(headers[j - 1]) + 4))
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def _ratio_total(df, res, c, base_m, col):
    num, den = MEASURES[base_m][3]
    cv = c.split("｜")[0] if "｜" in c else None
    if col and cv and cv != "合計":
        sub = df[df[col] == cv]
    else:
        sub = df   # 合計欄群或無欄維度 → 全體
    dn = float(sub[den].sum()) if den in sub.columns else 0.0
    nu = float(sub[num].sum()) if num in sub.columns else 0.0
    return (nu / dn) if dn else None
