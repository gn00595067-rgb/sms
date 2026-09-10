"""
render_excel.py — DataFrame → 格式化 Excel bytes（openpyxl）

規格（CLAUDE_CODE_TASK §9.3）：
  第1列標題（粗體14）、第2列篩選條件與產出時間、第4列表頭（底色/粗體/置中/凍結）、
  資料列（money #,##0、pct 0.0%、負數紅字）、最後合計列（粗體/上框線）、欄寬自動、微軟正黑體。
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.format import (DATE_COLS, INT_COLS, MONEY_COLS, PERCENT_NUM_COLS,
                         RATIO_PCT_COLS, YM_COLS, date_text, label, ym_text)

FONT_NAME = "微軟正黑體"
HEADER_FILL = PatternFill("solid", fgColor="1F5F8B")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
SUB_FONT = Font(name=FONT_NAME, size=9, color="666666")
BASE_FONT = Font(name=FONT_NAME)
BOLD_FONT = Font(name=FONT_NAME, bold=True)
RED_FONT = Font(name=FONT_NAME, color="C00000")
TOP_BORDER = Border(top=Side(style="thin"))


def _num_format(col: str) -> str | None:
    if col in MONEY_COLS or col in INT_COLS:
        return "#,##0"
    if col in RATIO_PCT_COLS:
        return "0.0%"
    if col in PERCENT_NUM_COLS:
        return '0.0"%"'
    return None


def build_excel(df: pd.DataFrame, title: str, filter_text: str = "",
                columns: list[str] | None = None, totals: dict | None = None) -> bytes:
    if df is None:
        df = pd.DataFrame()
    cols = [c for c in (columns or list(df.columns)) if c in df.columns] if not df.empty else (columns or [])
    wb = Workbook()
    ws = wb.active
    ws.title = "報表"
    ncol = max(len(cols), 1)

    # 第 1 列：標題
    ws.cell(1, 1, title).font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    # 第 2 列：篩選條件 + 產出時間
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    sub = (filter_text + "　" if filter_text else "") + f"產出時間：{stamp}"
    ws.cell(2, 1, sub).font = SUB_FONT
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)

    # 第 4 列：表頭
    header_row = 4
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(header_row, j, label(c))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 資料列
    r = header_row + 1
    for _, row in df.iterrows() if not df.empty else []:
        for j, c in enumerate(cols, start=1):
            v = row[c]
            fmt = _num_format(c)
            if c in YM_COLS:
                v = ym_text(v)
            elif c in DATE_COLS:
                v = date_text(v)
            elif fmt is not None and v is not None and not (isinstance(v, float) and pd.isna(v)):
                try:
                    v = float(v)
                except (ValueError, TypeError):
                    fmt = None
            cell = ws.cell(r, j, v if not (isinstance(v, float) and pd.isna(v)) else None)
            cell.font = RED_FONT if (fmt and isinstance(v, (int, float)) and v < 0) else BASE_FONT
            if fmt:
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")
        r += 1

    # 合計列
    if totals:
        first = True
        for j, c in enumerate(cols, start=1):
            v = totals.get(c)
            cell = ws.cell(r, j)
            if first:
                cell.value = "合計"
                first = False
            if v is not None:
                fmt = _num_format(c)
                cell.value = float(v) if fmt else v
                if fmt:
                    cell.number_format = fmt
                    cell.alignment = Alignment(horizontal="right")
            cell.font = BOLD_FONT
            cell.border = TOP_BORDER
        r += 1

    # 凍結表頭 + 欄寬
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    for j, c in enumerate(cols, start=1):
        width = len(label(c)) * 2 + 2
        if not df.empty:
            lengths = df[c].head(50).map(lambda v: 0 if v is None or (isinstance(v, float) and pd.isna(v)) else len(str(v)))
            width = max(width, int(lengths.max() or 0) + 2)
        ws.column_dimensions[get_column_letter(j)].width = min(max(width, 8), 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
