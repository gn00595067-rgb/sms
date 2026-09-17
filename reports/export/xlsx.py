"""
xlsx.py — Doc → Excel（openpyxl）

版面：
  「摘要」工作表：標題、期間/口徑/篩選、KPI 表、說明、目錄（超連結到各表）
  每個 Table 一個工作表：標題列、副標、（兩層）表頭、資料、合計、註記；凍結、篩選、欄寬、列印設定
  「圖表」工作表：有 excel 提示的圖以原生 Excel 圖（長條/堆疊/折線）呈現，資料在旁邊，可再編輯
列印設定：A4、方向依 Doc、縮放至一頁寬、每頁重複表頭、頁首頁尾（標題/期間/頁碼）。
"""
from __future__ import annotations

import io
import math
import re
import unicodedata
from datetime import date, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

from .doc import Chart, Col, Doc, Heading, Kpis, Row, Table, Text
from .theme import (COMPANY, NUM_FMT, PLATFORM, X_BASE, X_BOLD, X_CELL_BORDER, X_CENTER, X_GRP_FILL,
                    X_HEAD, X_HEAD_BORDER, X_HEAD_FILL, X_KPI_FILL, X_LEFT, X_MUTED, X_RIGHT, X_SUB,
                    X_TITLE, X_TOTAL_BORDER, X_ZEBRA_FILL, XLSX_FONT)

_BAD = re.compile(r"[\[\]:*?/\\]")


def _missing(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    try:
        return bool(pd.isna(v)) if not isinstance(v, (list, tuple, dict, str)) else False
    except (TypeError, ValueError):
        return False


def _width(s) -> int:
    if s is None:
        return 0
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in str(s))


def _sheet_name(wb: Workbook, name: str) -> str:
    base = _BAD.sub("", name).strip() or "報表"
    base = base[:28]
    n, out = 1, base
    while out in wb.sheetnames:
        n += 1
        out = f"{base}{n}"
    return out


def _print_setup(ws, *, orientation: str, title: str, period: str, who: str, header_rows: str | None) -> None:
    ws.page_setup.orientation = "landscape" if orientation == "landscape" else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.7, bottom=0.7, header=0.3, footer=0.3)
    if header_rows:
        ws.print_title_rows = header_rows
    # 頁首頁尾：不要用 &8 這種字級碼——後面接數字（2025…）會被當成兩位數字級；字級交給 Excel 預設
    ws.oddHeader.left.text = f'&"{XLSX_FONT}"{title}'
    ws.oddHeader.right.text = f'&"{XLSX_FONT}"期間 {period}' if period else ""
    ws.oddFooter.left.text = f'&"{XLSX_FONT}"{who}'
    ws.oddFooter.right.text = f'&"{XLSX_FONT}"第 &P / &N 頁'
    ws.sheet_view.showGridLines = False


def _title_rows(ws, doc: Doc, title: str, ncol: int) -> int:
    """寫標題 + 副標；回傳下一個可用列。"""
    ws.cell(1, 1, title).font = X_TITLE
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncol, 1))
    parts = [p for p in (doc.title, doc.period_text, doc.scope_text, doc.filter_text) if p]
    sub = "　".join(parts) + f"　產出：{(doc.generated_by + ' ') if doc.generated_by else ''}{doc.generated_at:%Y/%m/%d %H:%M}"
    c = ws.cell(2, 1, sub)
    c.font = X_SUB
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncol, 1))
    return 4


def _write_value(cell, v, kind: str):
    if _missing(v):
        cell.value = None
        return
    if kind in ("money", "int", "pct", "progress"):
        try:
            cell.value = float(v)
            cell.number_format = NUM_FMT[kind]
            cell.alignment = X_RIGHT
            return
        except (TypeError, ValueError):
            pass
    if kind == "ym" and isinstance(v, (date, datetime)):
        cell.value = f"{v.year}/{v.month:02d}"
        cell.alignment = X_LEFT
        return
    if kind == "date" and isinstance(v, (date, datetime)):
        cell.value = v if isinstance(v, datetime) else datetime(v.year, v.month, v.day)
        cell.number_format = "yyyy/mm/dd"
        cell.alignment = X_LEFT
        return
    cell.value = str(v)
    cell.alignment = X_LEFT


def _table_sheet(wb: Workbook, doc: Doc, t: Table) -> str:
    name = _sheet_name(wb, t.sheet_name or t.title)
    ws = wb.create_sheet(name)
    df = t.df if t.df is not None else pd.DataFrame()
    cols = [c for c in t.columns if c.key in df.columns] if not df.empty else list(t.columns)
    ncol = len(cols)
    r = _title_rows(ws, doc, t.title, ncol)

    # 兩層表頭
    if t.col_groups:
        j = 1
        for label, n in t.col_groups:
            if label:
                ws.cell(r, j, label).font = X_HEAD
                ws.cell(r, j).fill = X_GRP_FILL
                ws.cell(r, j).alignment = X_CENTER
                if n > 1:
                    ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=j + n - 1)
            j += n
        r += 1
    hdr = r
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(hdr, j, c.label)
        cell.font, cell.fill, cell.border = X_HEAD, X_HEAD_FILL, X_HEAD_BORDER
        cell.alignment = X_CENTER
    ws.row_dimensions[hdr].height = 20

    # 資料
    r = hdr + 1
    first_data = r
    for i, (_, row) in enumerate(df.iterrows()):
        rc = t.row_class(row) if t.row_class else None
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(r, j)
            _write_value(cell, row[c.key], c.kind)
            cell.font = X_MUTED if rc == "muted" else X_BASE
            cell.border = X_CELL_BORDER
            if i % 2 == 1:
                cell.fill = X_ZEBRA_FILL
        r += 1
    last_data = r - 1

    # 合計
    if t.totals:
        has_label = any(c.kind == "text" and t.totals.get(c.key) is not None for c in cols)
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(r, j)
            v = t.totals.get(c.key)
            if j == 1 and (v is None or c.kind == "text"):
                cell.value = str(v) if v is not None else ("" if has_label else "合計")
            elif v is not None:
                _write_value(cell, v, "pct" if c.kind == "progress" else c.kind)
            cell.font = X_BOLD
            cell.border = X_TOTAL_BORDER
        r += 1
    if t.note:
        ws.cell(r + 1, 1, t.note).font = X_SUB
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=max(ncol, 1))
    ws.cell(r + 2 if not t.note else r + 2, 1, "單位：元；比率＝Σ分子÷Σ分母").font = X_SUB

    # progress → 資料橫條
    for j, c in enumerate(cols, start=1):
        if c.kind == "progress" and last_data >= first_data:
            rng = f"{get_column_letter(j)}{first_data}:{get_column_letter(j)}{last_data}"
            ws.conditional_formatting.add(rng, DataBarRule(start_type="num", start_value=0, end_type="max", color="9DC3EC"))

    # 凍結、篩選、欄寬
    ws.freeze_panes = ws.cell(hdr + 1, min(t.freeze_cols, ncol) + 1)
    if last_data >= first_data:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(ncol)}{last_data}"
    for j, c in enumerate(cols, start=1):
        w = _width(c.label) + 2
        tv = (t.totals or {}).get(c.key)
        if not df.empty:
            sample = df[c.key].head(300)
            if c.kind in ("money", "int"):
                nums = pd.to_numeric(pd.concat([sample, pd.Series([tv])]), errors="coerce").abs().max()
                w = max(w, len(f"{nums:,.0f}") + 4 if pd.notna(nums) else 12, 12)   # 粗體合計要多留 2–3 字寬
            elif c.kind in ("pct", "progress"):
                w = max(w, 9)
            else:
                w = max(w, int(sample.map(_width).max() or 0) + 2, _width(tv) + 2 if isinstance(tv, str) else 0)
        ws.column_dimensions[get_column_letter(j)].width = min(max(w, 7), 44)
    _print_setup(ws, orientation=doc.orientation if not t.wide else "landscape", title=f"{doc.title}｜{t.title}",
                 period=doc.period_text, who=f"產出：{doc.generated_by} {doc.generated_at:%Y/%m/%d %H:%M}",
                 header_rows=f"{hdr - 1 if t.col_groups else hdr}:{hdr}")
    return name


def _kpi_value(v):
    """KPI 顯示字串 → 儘量轉數值（'71,533,523' / '59.2%' / '116 家'）。"""
    s = str(v).strip()
    try:
        if s.endswith("%"):
            return float(s[:-1].replace(",", "")) / 100, NUM_FMT["pct"]
        return float(s.replace(",", "")), NUM_FMT["money"]
    except ValueError:
        return s, None


def _summary_sheet(wb: Workbook, doc: Doc, table_sheets: list[tuple[str, str]], chart_sheet: str | None) -> None:
    ws = wb.active
    ws.title = "摘要"
    ws.cell(1, 1, doc.title).font = X_TITLE
    r = 2
    if doc.subtitle:
        ws.cell(r, 1, doc.subtitle).font = X_SUB
        r += 1
    meta = [f"期間：{doc.period_text}" if doc.period_text else None,
            f"口徑：{doc.scope_text}" if doc.scope_text else None,
            f"篩選：{doc.filter_text}" if doc.filter_text else None,
            f"產出：{doc.generated_by + ' ' if doc.generated_by else ''}{doc.generated_at:%Y/%m/%d %H:%M}",
            doc.as_of_note]
    for m in [x for x in meta if x]:
        ws.cell(r, 1, m).font = X_SUB
        r += 1
    r += 1
    for b in doc.blocks:
        items = b.blocks if isinstance(b, Row) else [b]
        for x in items:
            if isinstance(x, Heading):
                ws.cell(r, 1, x.text).font = X_BOLD
                r += 1
            elif isinstance(x, Kpis):
                if x.title:
                    ws.cell(r, 1, x.title).font = X_BOLD
                    r += 1
                for j, h in enumerate(("指標", "數值", "同期", "說明"), start=1):
                    c = ws.cell(r, j, h)
                    c.font, c.fill, c.border, c.alignment = X_HEAD, X_HEAD_FILL, X_HEAD_BORDER, X_CENTER
                r += 1
                for it in x.items:
                    ws.cell(r, 1, str(it.get("label", ""))).font = X_BASE
                    val, fmt = _kpi_value(it.get("value", ""))
                    c = ws.cell(r, 2, val)
                    c.font, c.alignment = X_BOLD, X_RIGHT
                    if fmt:
                        c.number_format = fmt
                    d = it.get("delta")
                    if d is not None and not _missing(d):
                        if isinstance(d, str):
                            ws.cell(r, 3, d).font = X_BASE
                        else:
                            c3 = ws.cell(r, 3, float(d))
                            c3.number_format = '+0.0%;[Red]-0.0%;"="'
                            c3.font, c3.alignment = X_BASE, X_RIGHT
                    if it.get("sub"):
                        ws.cell(r, 4, str(it["sub"]).replace("**", "")).font = X_MUTED
                    for j in range(1, 5):
                        ws.cell(r, j).fill = X_KPI_FILL
                        ws.cell(r, j).border = X_CELL_BORDER
                    r += 1
                r += 1
            elif isinstance(x, Text):
                c = ws.cell(r, 1, x.text.replace("**", ""))
                c.font = X_SUB if x.style == "note" else X_BASE
                c.alignment = Alignment(wrap_text=True, vertical="top")
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
                ws.row_dimensions[r].height = 15 * (1 + x.text.count("\n") + _width(x.text) // 110)
                r += 1
    r += 1
    ws.cell(r, 1, "目錄").font = X_BOLD
    r += 1
    for title, sheet in table_sheets:
        c = ws.cell(r, 1, title)
        c.hyperlink = f"#'{sheet}'!A1"
        c.font = Font(name=XLSX_FONT, size=10, color="2A78D6", underline="single")
        r += 1
    if chart_sheet:
        c = ws.cell(r, 1, "圖表（原生 Excel 圖，可編輯）")
        c.hyperlink = f"#'{chart_sheet}'!A1"
        c.font = Font(name=XLSX_FONT, size=10, color="2A78D6", underline="single")
        r += 1
    else:
        ws.cell(r, 1, "圖表請見 PDF 版").font = X_SUB
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 40
    _print_setup(ws, orientation="portrait", title=doc.title, period=doc.period_text,
                 who=f"產出：{doc.generated_by} {doc.generated_at:%Y/%m/%d %H:%M}", header_rows=None)


def _chart_sheet(wb: Workbook, doc: Doc, charts: list[Chart]) -> str | None:
    items = [c for c in charts if c.excel and isinstance(c.excel.get("df"), pd.DataFrame) and not c.excel["df"].empty]
    if not items:
        return None
    ws = wb.create_sheet(_sheet_name(wb, "圖表"))
    _title_rows(ws, doc, "圖表", 8)
    r = 4
    for ch in items:
        spec = ch.excel
        df, x, series = spec["df"], spec["x"], [s for s in spec["series"] if s in spec["df"].columns]
        unit = spec.get("unit", "元")
        ws.cell(r, 1, ch.title).font = X_BOLD
        r += 1
        # 資料區
        ws.cell(r, 1, x).font = X_HEAD
        ws.cell(r, 1).fill = X_HEAD_FILL
        for j, s in enumerate(series, start=2):
            c = ws.cell(r, j, s)
            c.font, c.fill, c.alignment = X_HEAD, X_HEAD_FILL, X_CENTER
        top = r
        for _, row in df.iterrows():
            r += 1
            ws.cell(r, 1, str(row[x])).font = X_BASE
            for j, s in enumerate(series, start=2):
                v = row[s]
                c = ws.cell(r, j, None if _missing(v) else float(v))
                c.number_format = "#,##0" if unit != "%" else "0.0%"
                c.font = X_BASE
        bottom = r
        kind = spec.get("type", "bar")
        if kind == "matrix":   # 熱圖 → 數字矩陣 + 色階（Excel 沒有原生熱圖）
            rng = f"B{top + 1}:{get_column_letter(1 + len(series))}{bottom}"
            ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                              end_type="max", end_color="6353BC"))
            r = bottom + 3
            continue
        chart = LineChart() if kind == "line" else BarChart()
        if kind != "line":
            chart.type = "col"
            if kind == "stacked":
                chart.grouping = "stacked"
                chart.overlap = 100
        chart.title = ch.title if (not unit or unit in ch.title) else f"{ch.title}（{unit}）"
        chart.y_axis.numFmt = "#,##0" if unit != "%" else "0%"
        chart.y_axis.majorGridlines = None if kind == "line" else chart.y_axis.majorGridlines
        chart.height, chart.width = 9, 22
        chart.legend.position = "b"
        data = Reference(ws, min_col=2, max_col=1 + len(series), min_row=top, max_row=bottom)
        cats = Reference(ws, min_col=1, min_row=top + 1, max_row=bottom)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        palette = {**COMPANY, **PLATFORM}
        for s_obj, name in zip(chart.series, series):
            col = palette.get(name)
            if col:
                if kind == "line":
                    s_obj.graphicalProperties.line.solidFill = col.lstrip("#")
                    s_obj.graphicalProperties.line.width = 22000
                else:
                    s_obj.graphicalProperties.solidFill = col.lstrip("#")
                    s_obj.graphicalProperties.line.solidFill = col.lstrip("#")
        ws.add_chart(chart, f"{get_column_letter(len(series) + 3)}{top}")
        r = max(bottom, top + 18) + 3
    ws.column_dimensions["A"].width = 16
    for j in range(2, 12):
        ws.column_dimensions[get_column_letter(j)].width = 13
    _print_setup(ws, orientation="landscape", title=f"{doc.title}｜圖表", period=doc.period_text,
                 who=f"產出：{doc.generated_by} {doc.generated_at:%Y/%m/%d %H:%M}", header_rows=None)
    return ws.title


def _single_sheet(wb: Workbook, doc: Doc) -> None:
    """所有 Heading／Table 依序寫在同一張工作表（年度發稿明細：三段往下排）。"""
    ws = wb.active
    ws.title = _sheet_name(wb, doc.title)
    ws.cell(1, 1, doc.title).font = X_TITLE
    parts = [p for p in (doc.period_text, doc.scope_text, doc.filter_text) if p]
    sub = "　".join(parts) + f"　產出：{(doc.generated_by + ' ') if doc.generated_by else ''}{doc.generated_at:%Y/%m/%d %H:%M}"
    ws.cell(2, 1, sub).font = X_SUB
    r = 4
    first_hdr = None
    maxcol = 1
    for b in doc.blocks:
        if isinstance(b, Heading):
            ws.cell(r, 1, b.text).font = X_BOLD if b.level >= 3 else X_HEAD
            r += 1
        elif isinstance(b, Text):
            c = ws.cell(r, 1, b.text.replace("**", ""))
            c.font = X_SUB if b.style == "note" else X_BASE
            r += 1
        elif isinstance(b, Table):
            df = b.df if b.df is not None else pd.DataFrame()
            cols = [c for c in b.columns if c.key in df.columns] if not df.empty else list(b.columns)
            maxcol = max(maxcol, len(cols))
            if b.title:
                ws.cell(r, 1, b.title).font = X_BOLD
                r += 1
            hdr = r
            if first_hdr is None:
                first_hdr = hdr
            for j, c in enumerate(cols, start=1):
                cell = ws.cell(hdr, j, c.label)
                cell.font, cell.fill, cell.border, cell.alignment = X_HEAD, X_HEAD_FILL, X_HEAD_BORDER, X_CENTER
            r = hdr + 1
            for i, (_, row) in enumerate(df.iterrows()):
                rc = b.row_class(row) if b.row_class else None
                for j, c in enumerate(cols, start=1):
                    cell = ws.cell(r, j)
                    _write_value(cell, row[c.key], c.kind)
                    cell.font = X_MUTED if rc == "muted" else X_BASE
                    cell.border = X_CELL_BORDER
                    if i % 2 == 1:
                        cell.fill = X_ZEBRA_FILL
                r += 1
            if b.totals:
                has_label = any(c.kind == "text" and b.totals.get(c.key) is not None for c in cols)
                for j, c in enumerate(cols, start=1):
                    cell = ws.cell(r, j)
                    v = b.totals.get(c.key)
                    if j == 1 and (v is None or c.kind == "text"):
                        cell.value = str(v) if v is not None else ("" if has_label else "合計")
                    elif v is not None:
                        _write_value(cell, v, "pct" if c.kind == "progress" else c.kind)
                    cell.font, cell.border = X_BOLD, X_TOTAL_BORDER
                r += 1
            if b.note:
                ws.cell(r, 1, b.note).font = X_SUB
                r += 1
            r += 2   # 段落間空兩列
    for j in range(1, maxcol + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = ws.cell((first_hdr or 4) + 1, 1)
    _print_setup(ws, orientation="landscape", title=doc.title, period=doc.period_text,
                 who=f"產出：{doc.generated_by} {doc.generated_at:%Y/%m/%d %H:%M}",
                 header_rows=f"{first_hdr}:{first_hdr}" if first_hdr else None)


def to_xlsx(doc: Doc) -> bytes:
    wb = Workbook()
    if getattr(doc, "xlsx_layout", "sheets") == "single":
        _single_sheet(wb, doc)
    else:
        table_sheets = [(t.title, _table_sheet(wb, doc, t)) for t in doc.tables()]
        chart_sheet = _chart_sheet(wb, doc, doc.charts())
        _summary_sheet(wb, doc, table_sheets, chart_sheet)
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
