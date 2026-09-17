"""
theme.py — 匯出層的視覺常數：字型、顏色、列印 CSS、Excel 樣式

原則：畫面（Streamlit）、PDF、Excel 三處同一套顏色與字級層級；紅黃綠只用在警示。
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---- 顏色（與 core/format.COLORS 一致）----
COMPANY = {"聲活": "#2a78d6", "東吳": "#1baf7a", "鉑霖": "#eb6834", "瑞迪": "#4a3aa7"}
PLATFORM = {"全家企頻": "#4a3aa7", "萬家福": "#6353bc", "新鮮視": "#8377cf", "廣播": "#a49bdf",
            "健康視": "#c3bde8", "營運": "#b8c0cc", "其它": "#d5dae2", "其他": "#d5dae2"}
ALERT = {"red": "#d03b3b", "yellow": "#fab219", "green": "#0ca30c"}
INK, MUTED, LINE, HEAD, ZEBRA, ACCENT = "#1d2733", "#66717f", "#d9e0e8", "#eef2f7", "#f7f9fc", "#2a78d6"

# ---- 字型 ----
# HTML/PDF：Windows 有微軟正黑體；Linux（Streamlit Cloud）用 packages.txt 裝 fonts-noto-cjk。
FONT_STACK = '"Noto Sans CJK TC","Noto Sans TC","Microsoft JhengHei","微軟正黑體","PingFang TC",sans-serif'
# Excel：Windows 開啟時用微軟正黑體；沒有的機器 Excel 會自動替代。
XLSX_FONT = "微軟正黑體"

# ---- 版面尺寸（A4，96dpi）----
PAGE = {
    "landscape": {"content_px": 1030, "margin": {"top": "16mm", "bottom": "14mm", "left": "11mm", "right": "11mm"}},
    "portrait":  {"content_px": 700,  "margin": {"top": "16mm", "bottom": "14mm", "left": "13mm", "right": "13mm"}},
}


def print_css(orientation: str) -> str:
    w = PAGE[orientation]["content_px"]
    half = int((w - 16) / 2)
    return f"""
:root {{ --ink:{INK}; --muted:{MUTED}; --line:{LINE}; --head:{HEAD}; --zebra:{ZEBRA}; --accent:{ACCENT}; }}
* {{ box-sizing: border-box; }}
html, body {{ margin: 0; padding: 0; background: #fff; }}
body {{ font-family: {FONT_STACK}; color: var(--ink); font-size: 11.5px; line-height: 1.45;
       -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
.page {{ width: {w}px; margin: 0 auto; }}
@page {{ size: A4 {orientation}; margin: 14mm 11mm 14mm 11mm; }}
@media screen {{ body {{ background: #e9edf2; }} .page {{ background: #fff; padding: 28px 24px; margin: 20px auto; box-shadow: 0 2px 12px rgba(0,0,0,.12); }} }}
@media print {{ .noprint {{ display: none !important; }} }}

/* 封面列 */
.cover {{ border-bottom: 2px solid var(--ink); padding-bottom: 8px; margin-bottom: 14px; }}
.cover h1 {{ font-size: 20px; margin: 0 0 2px; letter-spacing: .5px; }}
.cover .sub {{ color: var(--muted); font-size: 12px; margin: 0; }}
.cover .meta {{ display: flex; flex-wrap: wrap; gap: 6px 18px; font-size: 11px; color: var(--muted); margin-top: 6px; }}
.cover .meta b {{ color: var(--ink); font-weight: 600; }}
.asof {{ display: inline-block; background: #fff6da; border: 1px solid #f2d27a; color: #7a5a00; border-radius: 4px; padding: 1px 8px; font-size: 11px; }}

/* 節 */
h2 {{ font-size: 14px; margin: 16px 0 6px; padding-left: 8px; border-left: 4px solid var(--accent); break-after: avoid; page-break-after: avoid; }}
h3 {{ font-size: 12.5px; margin: 12px 0 4px; color: var(--ink); break-after: avoid; page-break-after: avoid; }}
.block {{ break-inside: avoid; page-break-inside: avoid; margin-bottom: 12px; }}
.tblock {{ margin-bottom: 12px; }}   /* 表格可跨頁：列不切、表頭每頁重複、標題跟著表 */
.row2 {{ display: flex; gap: 16px; align-items: flex-start; break-inside: avoid; page-break-inside: avoid; margin-bottom: 12px; }}
.row2 > * {{ flex: 1 1 0; min-width: 0; }}
.pb {{ break-before: page; page-break-before: always; }}

/* KPI */
.kpis {{ display: grid; gap: 10px; }}
.kpi {{ border: 1px solid var(--line); border-radius: 6px; padding: 8px 12px; background: #fff; }}
.kpi .l {{ font-size: 11px; color: var(--muted); }}
.kpi .v {{ font-size: 19px; font-weight: 700; line-height: 1.25; font-variant-numeric: tabular-nums; }}
.kpi .d {{ font-size: 11px; margin-top: 2px; }}
.kpi .d.up {{ color: {ALERT['green']}; }} .kpi .d.down {{ color: {ALERT['red']}; }} .kpi .d.flat {{ color: var(--muted); }}
.kpi .s {{ font-size: 10.5px; color: var(--muted); margin-top: 2px; }}
.kpi.co {{ border-top: 3px solid var(--co, var(--accent)); }}

/* 表格 */
table.t {{ width: 100%; border-collapse: collapse; font-size: 11px; font-variant-numeric: tabular-nums; }}
table.t.wide {{ font-size: 10px; }}
table.t thead {{ display: table-header-group; }}
table.t th {{ background: var(--head); color: var(--ink); font-weight: 600; padding: 5px 7px; border-bottom: 1.5px solid #b9c4d2; text-align: right; white-space: nowrap; vertical-align: bottom; }}
table.t th.grp {{ text-align: center; border-bottom: 1px solid #b9c4d2; background: #e4eaf2; font-weight: 600; }}
table.t th.txt, table.t td.txt {{ text-align: left; }}
table.t td {{ padding: 4px 7px; border-bottom: 1px solid var(--line); text-align: right; white-space: nowrap; vertical-align: middle; }}
table.t td.txt {{ white-space: normal; max-width: 260px; }}
table.t tr.z td {{ background: var(--zebra); }}
table.t tbody.keep {{ break-inside: avoid; page-break-inside: avoid; }}
table.t tr {{ break-inside: avoid; page-break-inside: avoid; }}
table.t tr.total td {{ font-weight: 700; border-top: 2px solid #9aa7b8; border-bottom: none; background: #fff; }}
table.t tr.muted td {{ color: #8a94a1; }}
table.t tr.warn td:first-child {{ box-shadow: inset 3px 0 0 {ALERT['red']}; }}
td.neg {{ color: {ALERT['red']}; }}
.bar {{ display: inline-block; height: 8px; background: #cfe0f7; border-radius: 2px; vertical-align: middle; margin-right: 6px; }}
.tnote {{ color: var(--muted); font-size: 10px; margin-top: 4px; }}
.trunc {{ color: #7a5a00; font-size: 10px; margin-top: 4px; }}

/* 圖 */
.chart {{ border: 1px solid var(--line); border-radius: 6px; padding: 6px 8px 2px; background: #fff; }}
.chart .ct {{ font-size: 12px; font-weight: 600; margin: 0 0 2px; }}
.chart .cap {{ color: var(--muted); font-size: 10.5px; margin: 2px 0 4px; }}
.chart .plot {{ width: 100%; }}
.chart.full .plot {{ width: {w - 18}px; }}
.chart.half .plot {{ width: {half - 18}px; }}

/* 文字 */
p.body {{ margin: 4px 0 8px; font-size: 11.5px; }}
p.note {{ margin: 2px 0 8px; font-size: 10.5px; color: var(--muted); }}
p.callout {{ margin: 6px 0 10px; padding: 8px 12px; background: #eef5fd; border-left: 4px solid var(--accent); border-radius: 4px; }}

/* 頁尾（瀏覽器列印用；Playwright 用 footer_template）*/
.foot {{ margin-top: 16px; padding-top: 6px; border-top: 1px solid var(--line); color: var(--muted); font-size: 10px; }}
.toolbar {{ position: sticky; top: 0; z-index: 9; background: #1d2733; color: #fff; padding: 8px 16px; display: flex; gap: 12px; align-items: center; font-size: 13px; }}
.toolbar button {{ background: var(--accent); color: #fff; border: 0; border-radius: 4px; padding: 6px 14px; font-size: 13px; cursor: pointer; }}
"""


# ---- Excel 樣式 ----
X_TITLE = Font(name=XLSX_FONT, size=14, bold=True, color="1D2733")
X_SUB = Font(name=XLSX_FONT, size=9, color="66717F")
X_BASE = Font(name=XLSX_FONT, size=10, color="1D2733")
X_BOLD = Font(name=XLSX_FONT, size=10, bold=True, color="1D2733")
X_MUTED = Font(name=XLSX_FONT, size=10, color="8A94A1")
X_HEAD = Font(name=XLSX_FONT, size=10, bold=True, color="1D2733")
X_HEAD_FILL = PatternFill("solid", fgColor="EEF2F7")
X_GRP_FILL = PatternFill("solid", fgColor="E4EAF2")
X_ZEBRA_FILL = PatternFill("solid", fgColor="F7F9FC")
X_KPI_FILL = PatternFill("solid", fgColor="F3F6FA")
_thin = Side(style="thin", color="D9E0E8")
_med = Side(style="medium", color="9AA7B8")
X_HEAD_BORDER = Border(bottom=Side(style="thin", color="B9C4D2"))
X_CELL_BORDER = Border(bottom=_thin)
X_TOTAL_BORDER = Border(top=_med)
X_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
X_RIGHT = Alignment(horizontal="right", vertical="center")
X_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=False)

NUM_FMT = {
    "money": '#,##0;[Red]-#,##0;"-"',
    "int": '#,##0;[Red]-#,##0;"-"',
    "pct": '0.0%;[Red]-0.0%;"-"',
    "progress": "0.0%",
}
