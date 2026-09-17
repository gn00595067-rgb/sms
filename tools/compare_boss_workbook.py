"""逐格比對：重現的工作簿 vs 老闆原檔（data_only 值）。數字容差 0.5；字串去空白比對。"""
from __future__ import annotations
import sys
import openpyxl

def norm(v):
    if v is None: return None
    if isinstance(v, (int, float)): return round(float(v), 1)
    return str(v).strip()

def compare(orig_path, new_path, sheets, max_show=15):
    wo = openpyxl.load_workbook(orig_path, data_only=True)
    wn = openpyxl.load_workbook(new_path, data_only=True)
    ok_all = True
    for name in sheets:
        a, b = wo[name], wn[name]
        rows = max(a.max_row, b.max_row); cols = max(a.max_column, b.max_column)
        diffs = []; n = 0
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                va, vb = norm(a.cell(r, c).value), norm(b.cell(r, c).value)
                if va is None and vb is None: continue
                n += 1
                same = (va == vb) or (isinstance(va, float) and isinstance(vb, float) and abs(va - vb) <= 0.5)
                if not same:
                    diffs.append((a.cell(r, c).coordinate, va, vb))
        print(f"{name}: {n} cells, {len(diffs)} diffs, orig rows {a.max_row} / new rows {b.max_row}")
        for d in diffs[:max_show]:
            print("   ", d)
        ok_all &= not diffs
    return ok_all

if __name__ == "__main__":
    sheets = ["平台總計總覽", "客戶數統計與客戶排名", "聲活_年度發稿明細", "東吳_年度發稿明細", "鉑霖_年度發稿明細"]
    compare(sys.argv[1], sys.argv[2], sheets)
