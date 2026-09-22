"""
test_finance.py — 財務報表專區（CLAUDE_CODE_TASK_7）黃金測試

三層：
  1. 純 Python（不用 DB）：舊系統排序（Big5 筆畫）、格式、Grid → HTML/Excel 不會炸。
  2. DB 結構：v_finance_line / v_channel_prepay / 規則表存在、種子筆數、預付日規則（PDF 第 31 頁那一週）。
  3. 黃金數字：以 V109g（2026/09/09 版 .mdb）資料算出、與《舊業績系統現有功能》PDF 截圖逐格對過的合計。
     只在 DB 是同一份快照時跑（2026/08 = 325 線、2026/09 = 227 線），否則 skip（換了資料快照數字自然不同）。
"""
from __future__ import annotations

import math
import pathlib
from datetime import date

import pytest

from tests.conftest import requires_db

ROOT = pathlib.Path(__file__).resolve().parent.parent
_builtin_round = round


def round(x, n=None):  # noqa: A001 — 舊報表（Access）是四捨五入，Python 內建 round 是四捨六入五成雙；黃金數字照舊報表
    return _builtin_round(x, n) if n else int(math.floor(float(x) + 0.5))


# ---------------------------------------------------------------- 1. 純 Python
def test_zh_key_matches_windows_big5_order():
    from reports.finance.compute import zh_key
    assert sorted(["聲活組", "東吳組", "鉑霖組", "錄音室", "聲活-東", "聲活-鉑"], key=zh_key) == \
        ["東吳組", "鉑霖組", "錄音室", "聲活-東", "聲活-鉑", "聲活組"]
    assert sorted(["陳絜心", "江曼瑄", "林怡慧"], key=zh_key) == ["江曼瑄", "林怡慧", "陳絜心"]
    assert sorted(["歐瑋君", "許雅婷", "陳秀鈴"], key=zh_key) == ["許雅婷", "陳秀鈴", "歐瑋君"]


def test_grid_fmt_and_renderers():
    from reports.finance.grid import Cell, Grid, fmt, grid_html, to_html, to_xlsx
    assert fmt(1234567.4, "money") == "1,234,567"
    assert fmt(-2800, "money") == "-2,800"
    assert fmt(0.2197, "pct2") == "21.97%"
    assert fmt(0.4295, "pct") == "43%"
    assert fmt(None, "money") == ""
    g = Grid("t", widths=[10, 10])
    g.add([Cell("標題", span=2, cls="b big noborder")], "title")
    g.add([Cell("a", cls="b ctr", fill="yellow"), Cell(-5, "money")])
    html = grid_html(g)
    assert 'colspan="2"' in html and "background:#FFFF00" in html and 'class="num neg"' in html
    assert "<title>x</title>" in to_html([g], title="x")
    data = to_xlsx([g])
    assert data[:2] == b"PK"
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data))["t"]
    assert ws["A1"].value == "標題" and ws["B2"].value == -5.0


def test_page_registered_for_finance_and_exec_only():
    txt = (ROOT / "app.py").read_text(encoding="utf-8")
    line = next(l for l in txt.splitlines() if '"finance_reports"' in l and "pages_app/finance_reports.py" in l)
    assert "FINANCE" in line and "EXEC" in line and "SALES" not in line and "MEDIA" not in line
    assert '"finance_reports"' in txt.split("GROUPS")[1].split("財務")[1].split("]")[0]   # 在「財務」群組


# ---------------------------------------------------------------- 2. DB 結構
def _count(query_df, sql, params=None) -> int:
    return int(query_df(sql, params).iloc[0, 0])


@requires_db
def test_finance_views_and_seed_tables():
    from core.data import query_df
    assert _count(query_df, "select count(*) from v_finance_line") == _count(query_df, "select count(*) from deal_line")
    assert _count(query_df, "select count(*) from channel_payment_rule") == 40
    assert _count(query_df, "select count(*) from channel_rebate_rate where year between 2024 and 2026") == 30
    cols = set(query_df("select column_name from information_schema.columns where table_name='v_finance_line'")["column_name"])
    for c in ("is_production", "plat_z", "fin_platform", "channel_rebate_pct", "channel_rebate_amount", "channel_cash_discount_pct",
              "network_name", "air_period_text", "ym", "profit", "contract_term"):
        assert c in cols, c
    mc = query_df("select name, cash_discount_pct, network_name from media_channel where name in ('中廣','BEST989','GOLD城市')").set_index("name")
    assert float(mc.loc["中廣", "cash_discount_pct"]) == 0.04
    assert mc.loc["BEST989", "network_name"] == "好事" and mc.loc["GOLD城市", "network_name"] == "城市"


@requires_db
def test_prepay_rule_matches_pdf_week():
    """PDF 第 31 頁：2026/09/07~09/13 → NEWS98 1150919 90,667（09/09）、大眾 1150401-9-6 15,750（09/09）。"""
    from core.data import query_df
    df = query_df("select media_channel, contract_no, prepay_on, cost_amount, pay_type from v_channel_prepay "
                  "where prepay_on between %s and %s", (date(2026, 9, 7), date(2026, 9, 13)))
    rows = {(r["media_channel"], r["contract_no"]): r for _, r in df.iterrows()}
    r = rows[("NEWS98", "1150919")]
    assert r["prepay_on"] == date(2026, 9, 9) and float(r["cost_amount"]) == 90667 and r["pay_type"] == "次週三付"
    r = rows[("大眾", "1150401-9-6")]
    assert r["prepay_on"] == date(2026, 9, 9) and float(r["cost_amount"]) == 15750
    # 每月2付：上檔日 ≤ 14 → 當月 15 日；> 14 → 次月 1 日（Bravo / 環宇 …）
    df = query_df("select prepay_on, air_start from v_channel_prepay where pay_type='每月2付' limit 200")
    for _, r in df.iterrows():
        d = r["air_start"]
        exp = date(d.year, d.month, 15) if d.day <= 14 else (date(d.year + (d.month == 12), d.month % 12 + 1, 1))
        assert r["prepay_on"] == exp


# ---------------------------------------------------------------- 3. 黃金數字（V109g 快照）
def _v109g(query_df) -> bool:
    return (_count(query_df, "select count(*) from deal_line where perf_ym = date '2026-08-01'") == 325
            and _count(query_df, "select count(*) from deal_line where perf_ym = date '2026-09-01'") == 227)


@pytest.fixture(scope="module")
def raw08():
    from core.data import query_df
    from reports.finance import compute as C
    if not _v109g(query_df):
        pytest.skip("DB 不是 V109g 快照，黃金數字不適用")
    return C.prepare(query_df("select * from v_finance_line where perf_ym = %s order by contract_no, line_no", (date(2026, 8, 1),)))


@pytest.fixture(scope="module")
def raw09():
    from core.data import query_df
    from reports.finance import compute as C
    if not _v109g(query_df):
        pytest.skip("DB 不是 V109g 快照，黃金數字不適用")
    return C.prepare(query_df("select * from v_finance_line where perf_ym = %s order by contract_no, line_no", (date(2026, 9, 1),)))


@pytest.fixture(scope="module")
def raw08_prev():
    from core.data import query_df
    from reports.finance import compute as C
    return C.prepare(query_df("select * from v_finance_line where perf_ym = %s", (date(2025, 8, 1),)))


@requires_db
def test_cost_report_final_page_2026_08(raw08):
    """PDF 第 11 頁：(1) 企頻 6,415,097 / 5,788,997 / 3,302,504 / 2,486,492 / 42.95%；企頻細項 全家、家樂福逐格相同。"""
    from reports.finance import compute as C
    rep = C.cost_report(raw08)
    f = rep.final.set_index("label")
    assert [round(f.loc["(1) 企頻", k]) for k in ("gross", "net", "cost", "profit")] == [6415097, 5788997, 3302504, 2486492]
    assert round(f.loc["(1) 企頻", "margin"], 4) == 0.4295
    assert [round(f.loc["(6) 營運", k]) for k in ("gross", "net", "cost", "profit")] == [1232205, 1232205, 1141558, 90647]
    cp = rep.cp_detail.set_index("label")
    assert [round(cp.loc["(1) 全家企頻", k]) for k in ("gross", "net", "cost", "profit")] == [2855775, 2687475, 1685602, 1001873]
    assert [round(cp.loc["(2) 家樂福企頻", k]) for k in ("gross", "net", "cost", "profit")] == [3559321, 3101521, 1616902, 1484619]
    assert round(f.loc["(1)-(7) 合計", "gross"]) == round(raw08["gross_amount"].sum())      # 合計 = 全部線（製作費列不重複加）
    # 東吳組總業績方塊（PDF 第 10 頁）：全家企頻 869,815 / 815,815 / 582,760 / 233,055 / 28.57%
    g = {name: box for name, _, box in rep.groups}["東吳組"].set_index("label")
    assert [round(g.loc["全家企頻", k]) for k in ("gross", "net", "cost", "profit")] == [869815, 815815, 582760, 233055]
    assert round(g.loc["全家企頻", "margin"], 4) == 0.2857


@requires_db
def test_combined_analysis_2026_08(raw08):
    """PDF 第 15 頁：聲活列與媒體頻道列逐格相同（新鮮視欄因資料快照差 148,500，不比）。"""
    from reports.finance import compute as C
    res = C.combined_analysis(raw08)
    ch = res["channel"].set_index("label")
    assert [round(ch.loc["實收金額", c]) for c in ("全家", "家樂福", "健康視", "廣播", "營運")] == [2855775, 3559321, 1408642, 2734304, 1232205]
    assert [round(ch.loc["除佣實收", c]) for c in ("全家", "家樂福", "健康視", "廣播")] == [2687475, 3101521, 1253257, 2635304]
    assert round(ch.loc["毛利率", "廣播"], 2) == 0.61
    g = res["by_group"]["實收金額"].set_index("label")
    assert [round(g.loc["聲活", c]) for c in C.COMB_COLS] == [642333, 750000, 1392333, 637924, 850000, 2880257, 2295381, 0, 1232205, 6407843]
    assert round(g.loc["聲+東+鉑", "全家"]) == 1983813 and round(g.loc["聲+東+鉑", "廣播"]) == 2295381   # 不含錄音室
    assert round(g.loc["聲-東鉑", "全家"]) == 871962
    assert round(res["by_group"]["除佣實收"].set_index("label").loc["聲活", "健康視"]) == 724925


@requires_db
def test_period_compare_2026_08(raw08, raw08_prev):
    """PDF 第 16 頁：2025 同期 總計 17,476,027 / 毛利 10,037,525；企頻 10,554,064；全家便利商店、鉑霖行動行銷、香港商群邑 逐格相同。"""
    from reports.finance import compute as C
    res = C.period_compare(raw08, raw08_prev, "全部")
    p = res["platform"].set_index("label")
    assert round(p.loc["總計", "net_prev"]) == 17476027 and round(p.loc["總計", "profit_prev"]) == 10037525
    assert round(p.loc["企頻", "net_prev"]) == 10554064 and round(p.loc["新鮮視", "net_prev"]) == 2683534 and round(p.loc["其他(廣播)", "net_prev"]) == 4238429
    assert round(p.loc["企頻", "profit_prev"]) == 7221222
    c = res["customers"].set_index("customer")
    r = c.loc["全家便利商店股份有限公司"]
    assert [round(r[k]) for k in ("net_cur", "net_prev", "net_growth", "profit_cur", "profit_prev", "profit_growth")] == [1897667, 1443333, 454334, 1314836, 285237, 1029599]
    assert round(r["net_rate"], 2) == 0.31 and round(r["profit_rate"], 2) == 3.61 and r["sp_cur"] == "Company" and r["sp_prev"] == "Company"
    r = c.loc["鉑霖行動行銷"]
    assert [round(r[k]) for k in ("net_cur", "net_prev", "net_growth", "profit_cur", "profit_prev")] == [1375868, 286257, 1089611, 1375868, 271314]
    r = c.loc["香港商群邑-2008"]
    assert [round(r[k]) for k in ("net_cur", "net_prev", "net_growth", "profit_cur", "profit_prev", "profit_growth")] == [901025, 2347600, -1446575, -90128, 2121948, -2212076]
    assert round(r["profit_rate"], 2) == -1.04


@requires_db
def test_bonus_summary_2026_08(raw08):
    """PDF 第 19 頁：陳絜心、林怡慧、許雅婷、歐瑋君、戴致祥（錄音室）逐格相同。"""
    from reports.finance import compute as C
    df = C.bonus_summary(raw08)
    rows = df[df["kind"] == "row"].set_index(["group", "salesperson"])
    r = rows.loc[("東吳組", "陳絜心")]
    assert [round(r[k]) for k in ("全家企頻", "家樂福企頻", "新鮮視", "gross", "net", "prod_cost", "recognized")] == [387244, 1360800, 191333, 2331482, 1939377, 23380, 1915997]
    r = rows.loc[("東吳組", "林怡慧")]
    assert [round(r[k]) for k in ("全家企頻", "家樂福企頻", "gross", "net", "prod_cost", "recognized")] == [428571, 47619, 476190, 476190, 2800, 473390]
    r = rows.loc[("鉑霖組", "許雅婷")]
    assert [round(r[k]) for k in ("全家企頻", "新鮮視", "gross", "net", "prod_cost", "recognized")] == [471665, 1197932, 1669597, 1669597, 45100, 1624497]
    r = rows.loc[("鉑霖組", "歐瑋君")]
    assert [round(r[k]) for k in ("新鮮視", "健康視", "gross", "net", "prod_cost", "recognized")] == [30000, 320201, 368571, 350201, 2080, 348121]
    r = rows.loc[("錄音室", "戴致祥")]
    assert [round(r[k]) for k in ("廣播", "gross", "net", "prod_cost", "recognized")] == [438923, 438923, 438923, 51860, 387063]
    r = rows.loc[("東吳組", "陳絜心換")]
    assert round(r["prod_cost"]) == 4300 and round(r["recognized"]) == -4300
    tot = df[(df["kind"] == "total") & (df["salesperson"] == "陳絜心 合計")].iloc[0]
    assert round(tot["recognized"]) == 1915997


@requires_db
def test_bonus_summary_platform_cost_2026_08(raw08):
    """製作成本分平台版：每平台除佣與原獎金總表相同；7 個平台製作成本相加 = 該列 prod_cost；認定不變。"""
    from reports.finance import compute as C, grid as G, layout as L
    base = C.bonus_summary(raw08)
    pc = C.bonus_summary_platform_cost(raw08)
    b = base[base["kind"] == "row"].set_index(["group", "salesperson", "month"])
    p = pc[pc["kind"] == "row"].set_index(["group", "salesperson", "month"])
    assert list(b.index) == list(p.index)                       # 列對齊
    for i in b.index:                                            # 除佣欄逐平台一致 + 平台製作成本相加 = prod_cost
        for plat in C.FIN_PLATFORMS:
            assert round(p.loc[i, plat]) == round(b.loc[i, plat])
        assert round(sum(p.loc[i, f"{plat}__pc"] for plat in C.FIN_PLATFORMS)) == round(p.loc[i, "prod_cost"])
    r = p.loc[("東吳組", "陳絜心", "08")]                          # 陳絜心 2026/08 製作成本 = 23,380（見 test_bonus_summary）
    assert round(r["prod_cost"]) == 23380
    assert round(sum(r[f"{plat}__pc"] for plat in C.FIN_PLATFORMS)) == 23380
    grids = L.layout_bonus_summary_platform_cost(pc, ym_from="2026/08", ym_to="2026/08")   # 版面可渲染
    html = G.to_html(grids, title="x")
    assert "除佣實收（分平台）" in html and "製作成本（分平台）" in html
    assert G.to_xlsx(grids)[:2] == b"PK"


@requires_db
def test_sales_wave_stats_2026_09(raw09):
    """PDF 第 20 頁：陳絜心 客戶數 8 / 實收 3,190,273 / 除佣 2,965,473 / 毛利 1,438,051 / 企頻 17 波段；洪佳琪 7 / 12 波段 / 2,798,048 / 2,369,085 / 654,832 / 企頻 6。"""
    from reports.finance import compute as C
    top = C.sales_wave_stats(raw09)["top"].set_index("salesperson")
    r = top.loc["陳絜心"]
    assert int(r["customers"]) == 8 and [round(r[k]) for k in ("gross", "net", "profit")] == [3190273, 2965473, 1438051]
    assert int(r["企頻_waves"]) == 17 and round(r["企頻_gross"]) == 2778607 and round(r["企頻_net"]) == 2553807
    r = top.loc["洪佳琪"]
    assert int(r["customers"]) == 7 and int(r["waves"]) == 12
    assert [round(r[k]) for k in ("gross", "net", "profit")] == [2798048, 2369085, 654832]
    assert int(r["企頻_waves"]) == 6 and round(r["企頻_gross"]) == 1445000 and round(r["新鮮視_gross"]) == 274048 and round(r["廣播_gross"]) == 829000
    assert list(top.index[:2]) == ["陳絜心", "洪佳琪"]          # 依除佣實收排序


@requires_db
def test_achievement_2026_09(raw09):
    """PDF 第 23–24 頁：東吳組 總計 2,572,856 / 21,950 / 2,550,906；鉑霖組 1,718,303 / 11,200 / 1,707,103；許雅婷 1,204,017 / 11,200 / 1,192,817。"""
    from reports.finance import compute as C
    res = {g["group"]: g for g in C.achievement(raw09)}
    assert [round(res["東吳組"]["total"][k]) for k in ("gross", "prod_cost", "recognized")] == [2572856, 21950, 2550906]
    assert [round(res["鉑霖組"]["total"][k]) for k in ("gross", "prod_cost", "recognized")] == [1718303, 11200, 1707103]
    sps = {s["salesperson"]: s for s in res["鉑霖組"]["salespeople"]}
    assert [round(sps["許雅婷"]["total"][k]) for k in ("gross", "prod_cost", "recognized")] == [1204017, 11200, 1192817]
    assert round(sps["歐瑋君"]["total"]["gross"]) == 514286
    sps = {s["salesperson"]: s for s in res["東吳組"]["salespeople"]}
    cats = {c["category"]: c for c in sps["陳絜心"]["categories"]}
    assert [round(cats["開發直客"]["sub"][k]) for k in ("gross", "prod_cost", "recognized")] == [2154761, 21950, 2132811]
    assert list(res) [:3] == ["東吳組", "鉑霖組", "聲活-東"]    # Big5 筆畫順（同舊報表）


@requires_db
def test_recognition_table_2026_09(raw09):
    """PDF 第 17–18 頁：洪佳琪 廣播 829,000 / 746,500 / D 4,800 → 741,700；新鮮視 274,048 / 268,548 / 480 → 268,068；健康視 250,000 / 187,425；家樂福企頻 620,000 / 483,712。"""
    from reports.finance import compute as C
    sp = next(s for s in C.recognition_table(raw09) if s["salesperson"] == "洪佳琪")
    blocks = {b["platform"]: b["totals"] for b in sp["months"][0]["blocks"]}
    assert [round(blocks["廣播"][k]) for k in ("gross", "net", "prod_cost", "recognized")] == [829000, 746500, 4800, 741700]
    assert [round(blocks["新鮮視"][k]) for k in ("gross", "net", "prod_cost", "recognized")] == [274048, 268548, 480, 268068]
    assert [round(blocks["健康視"][k]) for k in ("gross", "net")] == [250000, 187425]
    assert [round(blocks["家樂福企頻"][k]) for k in ("gross", "net")] == [620000, 483712]


@requires_db
def test_purchase_request_1150622():
    """PDF 第 28 頁：1150622 兩列（29,575 / 0）、除佣實收 75,075、毛利 45,500、60.61%。"""
    from core.data import query_df
    from reports.finance import compute as C
    lines = C.prepare(query_df("select * from v_finance_line where contract_no = %s order by line_no", ("1150622",)))
    if lines.empty:
        pytest.skip("DB 沒有 1150622")
    pr = C.purchase_request(lines)
    assert [round(x) for x in pr["rows"]["actual"]] == [29575, 0]
    assert round(pr["totals"]["actual"]) == 29575 and round(pr["net"]) == 75075 and round(pr["profit"]) == 45500
    assert round(pr["margin"], 4) == 0.6061


@requires_db
def test_layouts_render_html_and_xlsx(raw08, raw08_prev):
    from reports.finance import compute as C, grid as G, layout as L
    period = dict(ym_from="2026/08", ym_to="2026/08")
    grids = []
    grids += L.layout_cost_report(C.cost_report(raw08), **period)
    grids += L.layout_media_volume(C.media_volume(raw08, "客戶版", "實收金額"), version="客戶版", sort_by="實收金額", **period)
    grids += L.layout_combined(C.combined_analysis(raw08), **period)
    grids += L.layout_period_compare(C.period_compare(raw08, raw08_prev, "全部"), year=2026, **period)
    grids += L.layout_recognition(C.recognition_table(raw08), **period)
    grids += L.layout_bonus_summary(C.bonus_summary(raw08), **period)
    grids += L.layout_wave_stats(C.sales_wave_stats(raw08), **period)
    grids += L.layout_achievement(C.achievement(raw08), **period)
    html = G.to_html(grids, title="x")
    for needle in ("月業績成本報表", "6,415,097", "月電台發稿量客戶版 依實收金額排序", "綜合成本毛利分析", "客戶同期比較表",
                   "月業績認定表", "月獎金計算總表", "1,915,997", "業務發稿統計報表", "月責任檔業績達成表"):
        assert needle in html, needle
    assert G.to_xlsx(grids)[:2] == b"PK"
