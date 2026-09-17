"""
test_boss_reports.py — 階段 L 老闆版報表驗收（需要 DB）
  §4.1 年度發稿明細（鉑霖 2025）、§4.2 平台總計總覽（八區塊合計 / 客戶數）。
"""
from __future__ import annotations

import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / "scripts"))

from conftest import requires_db  # noqa: E402


@requires_db
def test_annual_detail_bolin_2025():
    from reports import annual_detail as AD
    df = AD.load(2025, "鉑霖", "media")
    plats = AD._present_platforms(df)
    ctot = float(df["net_amount"].sum())
    s1 = AD.section1(df, ctot, plats)
    hsu = s1[s1["業務"] == "許雅婷"].iloc[0]
    assert abs(float(hsu["總計"]) - 11_431_326) <= 1
    assert abs(float(hsu["帳上毛利"]) - 4_949_150) <= 1
    assert int(hsu["customers"]) == 54
    assert abs(float(hsu["利率"]) - 0.433) <= 0.001
    # 合計列
    tot = s1[s1["業務"] == "合計"].iloc[0]
    assert abs(float(tot["總計"]) - 11_907_516) <= 1
    # 第三段末列（各平台小計 + 總計）
    tl = AD.totals_line(df, plats)
    assert abs(float(tl["實收"]) - 11_907_516) <= 1
    assert abs(float(tl["成本"]) - 6_932_238) <= 1
    assert abs(float(tl["帳上毛利"]) - 4_975_277) <= 1


@requires_db
def test_annual_detail_workbook_opens():
    from reports import annual_detail as AD
    from openpyxl import load_workbook
    data = AD.build_workbook(2025, ["聲活", "東吳", "鉑霖"], "media")
    wb = load_workbook(io.BytesIO(data))
    assert "鉑霖_2025" in wb.sheetnames


@requires_db
def test_platform_overview_blocks_2025():
    from reports import platform_overview as PO
    df = PO._base(2025, "media")
    b1 = PO._block(df, "net", False)     # 除佣・含廣播
    tot1 = float(b1[b1["公司"] == "三公司合計"]["合計"].iloc[0])
    assert abs(tot1 - 134_303_021) <= 1
    b5 = PO._block(df, "net", True)      # 除佣・不含廣播（自媒體）
    tot5 = float(b5[b5["公司"] == "三公司合計"]["合計"].iloc[0])
    assert abs(tot5 - 98_408_698) <= 1
    # 客戶數 含廣播 / 不含廣播
    cust_all = int(PO._block(df, "customers", False)[lambda d: d["公司"] == "三公司合計"]["合計"].iloc[0])
    cust_own = int(PO._block(df, "customers", True)[lambda d: d["公司"] == "三公司合計"]["合計"].iloc[0])
    assert cust_all == 265 and cust_own == 259


@requires_db
def test_platform_overview_workbook_opens():
    from reports import platform_overview as PO
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(PO.build_workbook(2025, "media")))
    assert wb.sheetnames
