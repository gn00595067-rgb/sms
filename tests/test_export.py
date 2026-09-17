"""reports.export 的基本測試：HTML / Excel 一定要過；PDF 有瀏覽器引擎才測。"""
from __future__ import annotations

import io

import pandas as pd
import plotly.graph_objects as go
import pytest
from openpyxl import load_workbook

from reports import export as X


def _doc() -> X.Doc:
    df = pd.DataFrame({
        "customer": ["全家便利商店股份有限公司", "喬商廣告股份有限公司", "民視文化事業股份有限公司"],
        "net": [18469046.0, 6421685.0, 4698814.0],
        "cost": [11775849.0, 487280.0, 3034715.0],
        "margin": [0.3624, 0.9241, 0.3542],
        "share": [0.138, 0.048, 0.035],
        "neg": [-5.0, 0.0, 3.0],
    })
    doc = X.Doc("測試報表", subtitle="副標", period_text="2025/01–2025/12", scope_text="發稿口徑",
                filter_text="公司=全部", generated_by="tester")
    doc.kpis([{"label": "聲活 除佣實收", "value": "71,533,523", "delta": 0.12, "sub": "毛利 59.2%"},
              {"label": "客戶數", "value": "265 家", "delta": "不可比"}])
    fig = go.Figure(go.Bar(x=["聲活", "東吳", "鉑霖"], y=[7153, 5086, 1190]))
    doc.chart("公司業績（萬）", fig, excel={"type": "bar", "df": pd.DataFrame({"公司": ["聲活", "東吳", "鉑霖"], "萬": [7153, 5086, 1190]}),
                                          "x": "公司", "series": ["萬"], "unit": "萬"})
    doc.table("客戶排名", df, X.cols(("customer", "客戶", "text"), ("net", "除佣實收", "money"), ("cost", "成本", "money"),
                                    ("margin", "利率", "pct"), ("share", "佔比", "progress", {"max": 1.0}), ("neg", "負數", "int")),
              totals={"customer": "合計", "net": 29589545.0, "cost": 15297844.0, "margin": 0.483, "share": 0.221},
              note="單位：元")
    return doc


def test_html_contains_content():
    html = X.to_html(_doc(), plotly_js="none")
    assert "測試報表" in html and "18,469,046" in html and "36.2%" in html
    assert 'class="neg"' in html                    # 負數紅字
    assert "<thead>" in html and 'class="total"' in html
    assert "cdn.plot.ly" not in html                # plotly_js='none'


def test_html_splits_wide_tables():
    df = pd.DataFrame({"k": ["a", "b"], **{f"c{i}": [1.0, 2.0] for i in range(20)}})
    doc = X.Doc("寬表")
    doc.table("寬", df, X.cols(("k", "鍵", "text"), *[(f"c{i}", f"欄{i}", "money") for i in range(20)]))
    html = X.to_html(doc, plotly_js="none")
    assert "（續 2）" in html and html.count("<table") == 2


def test_xlsx_structure_and_print_setup():
    wb = load_workbook(io.BytesIO(X.to_xlsx(_doc())))
    assert wb.sheetnames[0] == "摘要" and "客戶排名" in wb.sheetnames and "圖表" in wb.sheetnames
    ws = wb["客戶排名"]
    assert ws.print_title_rows == "$4:$4" and ws.freeze_panes == "B5"
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1
    assert ws["B5"].value == 18469046.0 and ws["B5"].number_format.startswith("#,##0")
    assert ws["D5"].number_format.startswith("0.0%")
    assert ws.oddFooter.right.text.endswith("第 &P / &N 頁")
    assert "&8" not in ws.oddHeader.right.text          # 字級碼 + 年份 = 82pt 的陷阱
    assert len(wb["圖表"]._charts) == 1


def test_sheet_name_sanitized():
    doc = X.Doc("t")
    doc.table("客戶/業務: 排名 [2025]?", pd.DataFrame({"a": [1]}), X.cols(("a", "A", "int")))
    wb = load_workbook(io.BytesIO(X.to_xlsx(doc)))
    assert all(ch not in wb.sheetnames[1] for ch in "[]:*?/\\") and len(wb.sheetnames[1]) <= 31


@pytest.mark.skipif(not X.pdf_available(), reason="沒有 playwright")
def test_pdf_generates():
    pdf = X.to_pdf(_doc())
    if pdf is None:
        pytest.skip("此環境沒有可用的瀏覽器引擎")
    assert pdf[:4] == b"%PDF" and len(pdf) > 20_000
